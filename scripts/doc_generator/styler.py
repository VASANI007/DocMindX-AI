import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- Color Palette ---
COLOR_PRIMARY = RGBColor(10, 37, 64)       # Deep Navy #0A2540
COLOR_SECONDARY = RGBColor(0, 119, 182)    # Cerulean Blue #0077B6
COLOR_ACCENT = RGBColor(193, 18, 31)       # Crimson Red #C1121F
COLOR_SUCCESS = RGBColor(4, 120, 87)       # Emerald #047857
COLOR_DARK = RGBColor(30, 41, 59)          # Dark Slate #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)      # Slate Gray #64748B

HEX_NAVY = "0A2540"
HEX_CERULEAN = "0077B6"
HEX_CRIMSON = "C1121F"
HEX_EMERALD = "047857"
HEX_BG_LIGHT = "F1F5F9"
HEX_BG_ALT = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "0077B6"

def create_base_document(title, subtitle, doc_type, lang="EN"):
    doc = docx.Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"DocMindX AI — {doc_type} [{lang}]")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("DocMindX AI Platform | Enterprise Clinical Health Intelligence | Confidential & Proprietary")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_MUTED

    # Document Header Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(12)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Segoe UI"
        run_sub.font.size = Pt(12)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Banner Table
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_widths = [Inches(1.5), Inches(1.8), Inches(1.5), Inches(1.7)]
    
    headers_meta = (
        ("Project", "DocMindX AI v2.0", "Author", "Daksh Vasani"),
        ("Track", "Smart Health & Supply Chain", "Status", "Production Ready / Verified")
    ) if lang == "EN" else (
        ("प्रोजेक्ट", "DocMindX AI v2.0", "लेखक / आर्किटेक्ट", "Daksh Vasani"),
        ("ट्रैक", "स्मार्ट हेल्थ एवं सप्लाई चेन", "स्थिति", "प्रोडक्शन रेडी / सत्यापित")
    )

    for r_idx, row_data in enumerate(headers_meta):
        for c_idx in range(4):
            cell = meta_table.cell(r_idx, c_idx)
            cell.width = meta_widths[c_idx]
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            val = row_data[c_idx]
            run = cp.add_run(val)
            run.font.name = "Segoe UI"
            run.font.size = Pt(9)
            if c_idx % 2 == 0:
                run.font.bold = True
                run.font.color.rgb = COLOR_PRIMARY
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BG_LIGHT}"/>')
                cell._tc.get_or_add_tcPr().append(shd)
            else:
                run.font.color.rgb = COLOR_DARK
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_BG_ALT}"/>')
                cell._tc.get_or_add_tcPr().append(shd)

    set_table_borders(meta_table, HEX_BORDER)
    
    # Divider space
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(12)
    p_div.paragraph_format.space_after = Pt(8)
    
    return doc

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_body_p(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Segoe UI"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_DARK
    run.font.italic = italic
    return p

def add_bullet_p(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Segoe UI"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    run = p.add_run(text)
    run.font.name = "Segoe UI"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_DARK
    return p

def add_callout(doc, text, title=None, alert_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    border_hex = HEX_CERULEAN if alert_type == "info" else (HEX_CRIMSON if alert_type == "danger" else HEX_EMERALD)
    bg_hex = "F0F7FF" if alert_type == "info" else ("FEF2F2" if alert_type == "danger" else "F0FDF4")
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    tcPr.append(shd)
    
    # Left border only
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.right_indent = Inches(0.1)
    
    if title:
        rt = p.add_run(f"[{title}]\n")
        rt.font.name = "Segoe UI"
        rt.font.size = Pt(10.5)
        rt.font.bold = True
        rt.font.color.rgb = COLOR_ACCENT if alert_type == "danger" else (COLOR_SECONDARY if alert_type == "info" else COLOR_SUCCESS)
        
    rtxt = p.add_run(text)
    rtxt.font.name = "Segoe UI"
    rtxt.font.size = Pt(9.5)
    rtxt.font.color.rgb = COLOR_DARK
    
    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    tcPr.append(shd)
    
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="{HEX_CERULEAN}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    
    rc = p.add_run(code_text)
    rc.font.name = "Consolas"
    rc.font.size = Pt(8.5)
    rc.font.color.rgb = RGBColor(30, 41, 59)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(4)
    p_sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Style Header
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        if col_widths and c_idx < len(col_widths):
            cell.width = Inches(col_widths[c_idx])
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_NAVY}"/>')
        tcPr.append(shd)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(h_text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Style Data Rows
    for r_idx, row_data in enumerate(rows):
        bg = HEX_BG_ALT if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(r_idx + 1, c_idx)
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            tcPr.append(shd)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(str(val))
            run.font.name = "Segoe UI"
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_DARK

    set_table_borders(tbl, HEX_BORDER)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(6)
    p_sp.paragraph_format.space_after = Pt(4)
    return tbl

def set_table_borders(table, color_hex):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
